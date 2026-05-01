from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


OUTPUT_PATH = Path("Rehearse_Test_Cases.xlsx")
SHEET_NAME = "Full Test Execution"
VERSION = "v1"

HEADERS = [
    "Test Case ID",
    "Version",
    "Feature",
    "Test Case",
    "Preconditions",
    "Steps to Execute",
    "Test Data",
    "Expected Output",
    "Actual Output(TBD)",
    "Status(PENDING)",
    "Bug ID(-)",
    "Bug Title(-)",
    "Root Cause(-)",
    "Resolution(-)",
]

ROWS = [
    {
        "Test Case ID": "RH-TC-001",
        "Feature": "Authentication",
        "Test Case": "Sign up with valid email completes OTP verification and lands on onboarding",
        "Preconditions": "App launched; no active user session; reachable auth service.",
        "Steps to Execute": "Open auth screen. Choose Sign Up. Enter valid name, email, and password. Submit. Enter valid OTP from email.",
        "Test Data": "newuser@example.com / Valid 8+ char password / valid OTP",
        "Expected Output": "Account is created, OTP is accepted, session becomes authenticated, and onboarding flow opens without duplicate requests.",
    },
    {
        "Test Case ID": "RH-TC-002",
        "Feature": "Authentication",
        "Test Case": "Sign in with valid existing credentials opens the main tab bar",
        "Preconditions": "Existing verified account is present.",
        "Steps to Execute": "Open auth screen. Choose Log In. Enter valid email and password. Submit.",
        "Test Data": "existing@example.com / correct password",
        "Expected Output": "User is authenticated and routed to the home experience with profile-aware content.",
    },
    {
        "Test Case ID": "RH-TC-003",
        "Feature": "Authentication",
        "Test Case": "Invalid or malformed JWT/session prevents protected API access",
        "Preconditions": "Backend running with auth enabled.",
        "Steps to Execute": "Call protected endpoints such as /jobs with expired, malformed, missing, and wrong-secret tokens.",
        "Test Data": "Expired JWT / malformed bearer token / wrong signature JWT",
        "Expected Output": "Backend returns 401 or 403 and no protected data is exposed.",
    },
    {
        "Test Case ID": "RH-TC-004",
        "Feature": "Password Recovery",
        "Test Case": "Password recovery request sends OTP for a valid email",
        "Preconditions": "Existing account with recoverable email.",
        "Steps to Execute": "Open forgot-password flow. Enter valid account email. Submit request.",
        "Test Data": "user@example.com",
        "Expected Output": "Recovery OTP request succeeds and user is taken to the verification step.",
    },
    {
        "Test Case ID": "RH-TC-005",
        "Feature": "Password Recovery",
        "Test Case": "Password reset rejects mismatched new password and confirmation",
        "Preconditions": "Recovery flow is open.",
        "Steps to Execute": "Enter a new password and a different confirmation password. Submit.",
        "Test Data": "password123 / password124",
        "Expected Output": "Validation error is shown, password is not updated, and user stays on the reset screen.",
    },
    {
        "Test Case ID": "RH-TC-006",
        "Feature": "OTP Verification",
        "Test Case": "Pasting a full 6-digit OTP distributes digits across all inputs",
        "Preconditions": "OTP verification screen is visible.",
        "Steps to Execute": "Paste a 6-digit code into the OTP entry area.",
        "Test Data": "123456",
        "Expected Output": "All six fields are populated in order and focus advances to the last digit.",
    },
    {
        "Test Case ID": "RH-TC-007",
        "Feature": "Guest Session",
        "Test Case": "First guest launch shows info card and creates guest identity",
        "Preconditions": "Fresh install or cleared local guest state; no authenticated user.",
        "Steps to Execute": "Launch the app without signing in.",
        "Test Data": "New device state",
        "Expected Output": "Guest ID is created, info card is shown once, and guest mode continues into the app.",
    },
    {
        "Test Case ID": "RH-TC-008",
        "Feature": "Guest Access Control",
        "Test Case": "Guest user is blocked from sheet upload and prompted to sign up or log in",
        "Preconditions": "App is in guest mode; upload screen is accessible.",
        "Steps to Execute": "Open Upload. Try starting the scanner or import flow.",
        "Test Data": "Guest session",
        "Expected Output": "Guest feature gate modal appears and upload does not proceed until auth is chosen.",
    },
    {
        "Test Case ID": "RH-TC-009",
        "Feature": "Guest Access Control",
        "Test Case": "Guest user is blocked from chord recognition",
        "Preconditions": "App is in guest mode.",
        "Steps to Execute": "Navigate to Chord Recognition.",
        "Test Data": "Guest session",
        "Expected Output": "Screen redirects out, guest gate modal is presented, and microphone flow does not start.",
    },
    {
        "Test Case ID": "RH-TC-010",
        "Feature": "Onboarding",
        "Test Case": "Authenticated onboarding saves level, genres, practice goal, and avatar seed",
        "Preconditions": "Authenticated user just completed sign up.",
        "Steps to Execute": "Complete all onboarding questions and save.",
        "Test Data": "Level=Intermediate; Genres=Classical,Jazz & Blues; Practice=20 mins",
        "Expected Output": "Profile and user_onboarding are upserted and app transitions to the home tab bar.",
    },
    {
        "Test Case ID": "RH-TC-011",
        "Feature": "Onboarding",
        "Test Case": "Skipping onboarding still navigates to home with default practice goal",
        "Preconditions": "New guest or authenticated user in onboarding.",
        "Steps to Execute": "Use skip action before finishing all selections.",
        "Test Data": "No explicit selections or partial selections",
        "Expected Output": "User is not blocked, defaults are stored where applicable, and home screen opens.",
    },
    {
        "Test Case ID": "RH-TC-012",
        "Feature": "Home",
        "Test Case": "Home hero card prioritizes most recent play over recommendations",
        "Preconditions": "Authenticated user has at least one recent play.",
        "Steps to Execute": "Open Home after playing a song.",
        "Test Data": "Recent play exists for current user",
        "Expected Output": "Hero/continue card shows the most recent song instead of fallback recommendations.",
    },
    {
        "Test Case ID": "RH-TC-013",
        "Feature": "Home",
        "Test Case": "New user with no recents gets recommendation from onboarding level",
        "Preconditions": "Authenticated user has no recents; onboarding level exists.",
        "Steps to Execute": "Open Home.",
        "Test Data": "user_onboarding.level=some_basics",
        "Expected Output": "Home fetches the matching level recommendation and displays it as the top card.",
    },
    {
        "Test Case ID": "RH-TC-014",
        "Feature": "Home",
        "Test Case": "If no level match exists, home falls back to discover songs or hides hero card",
        "Preconditions": "No recents and no matching onboarding recommendation.",
        "Steps to Execute": "Open Home with empty recommendation set.",
        "Test Data": "No songs at target level; discover returns 0 or 1+ songs",
        "Expected Output": "First discover song is shown if available; otherwise hero card is hidden cleanly.",
    },
    {
        "Test Case ID": "RH-TC-015",
        "Feature": "Daily Goal",
        "Test Case": "Daily goal refresh updates UI when app reappears or auth state changes",
        "Preconditions": "Goal exists in Supabase or guest onboarding data.",
        "Steps to Execute": "Open Home. Background and foreground app or sign in after guest use.",
        "Test Data": "Practice goal of 10-30 minutes",
        "Expected Output": "Progress and goal labels refresh without duplicate crashes or stale values.",
    },
    {
        "Test Case ID": "RH-TC-016",
        "Feature": "Upload",
        "Test Case": "Upload a valid PDF creates a conversion job and shows authenticated result route",
        "Preconditions": "Authenticated user; backend and storage available.",
        "Steps to Execute": "Open Upload. Import a valid PDF. Complete upload flow.",
        "Test Data": "Well-formed PDF under file-size limit",
        "Expected Output": "Job is created with a job ID, status enters processing pipeline, and API exposes authenticated sheet routes instead of public URLs.",
    },
    {
        "Test Case ID": "RH-TC-017",
        "Feature": "Upload",
        "Test Case": "Upload a valid JPEG image is accepted for conversion",
        "Preconditions": "Authenticated user.",
        "Steps to Execute": "Open Upload. Import a valid JPEG. Submit.",
        "Test Data": "Well-formed JPEG with correct magic bytes",
        "Expected Output": "Backend accepts the file and returns a job ID for conversion.",
    },
    {
        "Test Case ID": "RH-TC-018",
        "Feature": "Upload Validation",
        "Test Case": "Disguised or corrupted files are rejected before processing",
        "Preconditions": "Authenticated user; upload endpoint reachable.",
        "Steps to Execute": "Attempt uploads for HTML renamed as PDF, random bytes with PDF type, PDF bytes labeled as JPEG, JPEG bytes labeled as PDF, empty file, and truncated PDF/JPEG.",
        "Test Data": "Malformed files with mismatched magic bytes and empty/truncated payloads",
        "Expected Output": "Endpoint returns validation failure such as 422 and no job is queued.",
    },
    {
        "Test Case ID": "RH-TC-019",
        "Feature": "Upload Validation",
        "Test Case": "Oversized upload is blocked by max file-size limit",
        "Preconditions": "Authenticated user.",
        "Steps to Execute": "Submit a file above configured MAX_FILE_SIZE.",
        "Test Data": "PDF larger than backend limit",
        "Expected Output": "Upload is rejected with 413 or validation error and nothing is stored for processing.",
    },
    {
        "Test Case ID": "RH-TC-020",
        "Feature": "Queue Management",
        "Test Case": "Upload is rejected gracefully when queue depth is at capacity",
        "Preconditions": "Queue depth equals MAX_QUEUE_DEPTH.",
        "Steps to Execute": "Submit a valid upload while the queue is full.",
        "Test Data": "Valid PDF; saturated Redis/RQ queue",
        "Expected Output": "API returns 503 capacity message, job is marked failed instead of lingering pending, and user is asked to retry later.",
    },
    {
        "Test Case ID": "RH-TC-021",
        "Feature": "Job Recovery",
        "Test Case": "Recovery loop requeues stuck processing jobs missing from StartedJobRegistry",
        "Preconditions": "Job exists in DB with status processing; worker entry expired or missing.",
        "Steps to Execute": "Run recovery cycle or wait for scheduled recovery loop.",
        "Test Data": "Processing job absent from StartedJobRegistry",
        "Expected Output": "Job status is reset to pending and the job is safely re-enqueued once.",
    },
    {
        "Test Case ID": "RH-TC-022",
        "Feature": "Backend Security",
        "Test Case": "User can only view their own jobs list and detail",
        "Preconditions": "Two separate users each have jobs.",
        "Steps to Execute": "Query /jobs and /jobs/{id} with each user token.",
        "Test Data": "User A token; User B token; mixed job IDs",
        "Expected Output": "Each user only receives their own jobs and cannot access the other user's details.",
    },
    {
        "Test Case ID": "RH-TC-023",
        "Feature": "Backend Security",
        "Test Case": "Authenticated sheet JSON and PDF routes deny cross-user access",
        "Preconditions": "User A has completed outputs; User B is authenticated separately.",
        "Steps to Execute": "Request /sheets/{job_id} and /sheets/{job_id}/pdf for another user's job.",
        "Test Data": "User B token with User A job ID",
        "Expected Output": "Backend returns 403 or 404 and never exposes the other user's asset.",
    },
    {
        "Test Case ID": "RH-TC-024",
        "Feature": "Signed URLs",
        "Test Case": "Signed URLs expire and public Supabase object URLs are never leaked in API responses",
        "Preconditions": "Completed job exists; storage signing configured.",
        "Steps to Execute": "Fetch convert response, job detail, and jobs list. Inspect returned URLs. Generate short-lived signed URL and wait past expiration.",
        "Test Data": "Completed output.json and labeled PDF paths",
        "Expected Output": "Responses use authenticated app routes, not public object/public URLs, and signed URLs stop working after expiry.",
    },
    {
        "Test Case ID": "RH-TC-025",
        "Feature": "Admin Orphan Cleanup",
        "Test Case": "Non-admin users are denied orphan detect and cleanup endpoints",
        "Preconditions": "Authenticated non-admin user.",
        "Steps to Execute": "Call /admin/orphans/detect and /admin/orphans/cleanup.",
        "Test Data": "Valid non-admin JWT",
        "Expected Output": "Endpoints return 401 or 403 and no cleanup action runs.",
    },
    {
        "Test Case ID": "RH-TC-026",
        "Feature": "Admin Orphan Cleanup",
        "Test Case": "Admin endpoint failures do not leak stack traces or storage/database internals",
        "Preconditions": "Authenticated admin user; orphan detector configured to raise server-side error.",
        "Steps to Execute": "Call admin orphan detection while backend throws an exception.",
        "Test Data": "Injected runtime failure text with file paths and table names",
        "Expected Output": "Client receives generic internal-server-error messaging while detailed error is kept server-side.",
    },
    {
        "Test Case ID": "RH-TC-027",
        "Feature": "Rate Limiting",
        "Test Case": "API rate limiter blocks repeated requests after configured limit",
        "Preconditions": "Rate limiting enabled.",
        "Steps to Execute": "Hit a protected or probe endpoint repeatedly above the configured threshold.",
        "Test Data": "Burst of requests above 2/minute or 5/minute test limit",
        "Expected Output": "Initial requests succeed, later requests return 429, and limiter does not fail open.",
    },
    {
        "Test Case ID": "RH-TC-028",
        "Feature": "Discover",
        "Test Case": "Search and filter by level and skill update the song list correctly",
        "Preconditions": "Discover song catalog exists.",
        "Steps to Execute": "Open Discover. Search by song/composer. Apply level filter. Apply skill filter. Clear chips.",
        "Test Data": "Mixed catalog with multiple levels and skill tags",
        "Expected Output": "Filtered list matches selected chips and search text; removing filters restores the correct set.",
    },
    {
        "Test Case ID": "RH-TC-029",
        "Feature": "Discover",
        "Test Case": "Empty-state messaging appears when no songs match active filters",
        "Preconditions": "Discover screen loaded.",
        "Steps to Execute": "Apply a search/filter combination that returns zero songs.",
        "Test Data": "Nonexistent term or unmatched level+skill combination",
        "Expected Output": "Table hides content cleanly and empty-state label displays without layout issues.",
    },
    {
        "Test Case ID": "RH-TC-030",
        "Feature": "Playback Animation",
        "Test Case": "Landscape playback loads sheet JSON, starts playback, and keeps progress in sync after tempo change and seek",
        "Preconditions": "Valid sheet music JSON or demo song is available.",
        "Steps to Execute": "Open animation screen. Start playback. Change tempo. Seek forward and backward. Pause and resume.",
        "Test Data": "Playable sheet JSON with multiple chords",
        "Expected Output": "Landscape UI renders correctly, piano and sheet stay synced, progress reflects tempo multiplier, and backward seek recomputes elapsed time correctly.",
    },
    {
        "Test Case ID": "RH-TC-031",
        "Feature": "Playback Animation",
        "Test Case": "Leaving playback returns app to portrait and ends tracked practice session cleanly",
        "Preconditions": "Animation screen is active and playing.",
        "Steps to Execute": "Open animation screen, begin playback, then navigate back.",
        "Test Data": "Any playable song/session",
        "Expected Output": "Playback resources are torn down, tracked session ends, tab bar returns, and app orientation resets to portrait.",
    },
    {
        "Test Case ID": "RH-TC-032",
        "Feature": "Chord Recognition",
        "Test Case": "Microphone-based note detection updates note, status, frequency, and waveform in real time",
        "Preconditions": "Authenticated non-guest user on device with microphone permission granted.",
        "Steps to Execute": "Open Chord Recognition. Start listening. Play single notes or chords into the microphone. Stop listening.",
        "Test Data": "Stable piano notes around C4-A4",
        "Expected Output": "Detected note and frequency update in near real time, waveform animates, and stop action halts listening cleanly.",
    },
    {
        "Test Case ID": "RH-TC-033",
        "Feature": "Playlists",
        "Test Case": "Create, open, and delete playlists including empty-state and cached-state behavior",
        "Preconditions": "Authenticated user with playlist permissions.",
        "Steps to Execute": "Open Playlists. Create a new playlist. Reopen screen. Open playlist detail. Delete single and multiple playlists.",
        "Test Data": "Playlist title, optional tags, optional cover",
        "Expected Output": "Playlists load from cache then remote, creation succeeds, empty-state updates correctly, and deletions confirm then remove items.",
    },
    {
        "Test Case ID": "RH-TC-034",
        "Feature": "Profile",
        "Test Case": "Profile screen loads stats, weekly progress, legal links, and guest-vs-auth menus correctly",
        "Preconditions": "Profile screen accessible.",
        "Steps to Execute": "Open Profile as guest and as authenticated user. Review stats, goal badge, menu, and legal actions.",
        "Test Data": "Guest onboarding data and authenticated lesson/practice data",
        "Expected Output": "Correct identity, stats, goal, and menu actions appear for the current session type; terms/privacy views open successfully.",
    },
    {
        "Test Case ID": "RH-TC-035",
        "Feature": "Profile Settings",
        "Test Case": "Changing theme updates appearance preference and persists across relaunch",
        "Preconditions": "Profile screen accessible.",
        "Steps to Execute": "Open Profile. Change Theme from current selection to Dark, Light, or System Default. Relaunch app.",
        "Test Data": "Each AppTheme option",
        "Expected Output": "Selected theme is applied immediately, stored in UserDefaults, notification is posted, and choice persists after relaunch.",
    },
    {
        "Test Case ID": "RH-TC-036",
        "Feature": "Profile Account Management",
        "Test Case": "Sign out or guest sign-out clears session state and prevents stale authenticated UI",
        "Preconditions": "User is signed in or using guest mode.",
        "Steps to Execute": "Open Profile. Tap sign out. Confirm action.",
        "Test Data": "Authenticated session and guest session",
        "Expected Output": "Current session is cleared, guest state is reset when applicable, and the app returns to the correct unauthenticated or guest entry state.",
    },
    {
        "Test Case ID": "RH-TC-037",
        "Feature": "Profile Account Management",
        "Test Case": "Delete account flow removes user account and related data without leaving access behind",
        "Preconditions": "Authenticated user; backend delete-account function available.",
        "Steps to Execute": "Open Profile menu. Choose delete account. Confirm destructive action and complete auth if required.",
        "Test Data": "Authenticated account with uploads, playlists, and profile data",
        "Expected Output": "Account deletion completes, user data is removed or queued for deletion per backend policy, and deleted account can no longer access protected endpoints.",
    },
    {
        "Test Case ID": "RH-TC-038",
        "Feature": "Performance",
        "Test Case": "Large song/discover lists and recent uploads remain responsive during load and refresh",
        "Preconditions": "Catalog, recents, or playlist datasets contain high item counts.",
        "Steps to Execute": "Open Home, Discover, Upload recents, and Playlists with large datasets. Scroll and revisit screens repeatedly.",
        "Test Data": "50-200 songs/uploads/playlists",
        "Expected Output": "UI remains responsive, loaders/empty states behave correctly, and repeated refreshes do not duplicate rows or freeze scrolling.",
    },
    {
        "Test Case ID": "RH-TC-039",
        "Feature": "Storage Resolution",
        "Test Case": "Asset resolver signs relative storage paths and preserves absolute Supabase URLs without double signing",
        "Preconditions": "Resolver configured with Supabase base URL and signer.",
        "Steps to Execute": "Resolve cover, PDF, and JSON assets using relative paths, legacy paths, and absolute public storage URLs.",
        "Test Data": "cover/user-1/song.png; user-1/file.pdf; https://demo.supabase.co/storage/v1/object/public/...",
        "Expected Output": "Relative paths are signed with the correct fallback bucket, legacy paths normalize correctly, and absolute Supabase URLs are returned directly without extra signing.",
    },
    {
        "Test Case ID": "RH-TC-040",
        "Feature": "Startup Configuration",
        "Test Case": "App and worker reject unsafe Audiveris or missing configuration on startup",
        "Preconditions": "Backend import/startup environment can be varied.",
        "Steps to Execute": "Start app or worker with remote HTTP Audiveris URL, missing URL, and approved localhost/HTTPS settings.",
        "Test Data": "AUDIVERIS_API_URL=http://remote-server.com/api and valid/invalid variants",
        "Expected Output": "Unsafe or missing configurations fail fast with clear startup errors; valid secure/local configurations start successfully.",
    },
]


def autofit_columns(ws) -> None:
    for column_cells in ws.columns:
        max_length = 0
        column_letter = get_column_letter(column_cells[0].column)
        for cell in column_cells:
            value = "" if cell.value is None else str(cell.value)
            max_length = max(max_length, len(value))
        ws.column_dimensions[column_letter].width = min(max_length + 2, 80)


def main() -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = SHEET_NAME

    header_fill = PatternFill(fill_type="solid", fgColor="1F4E78")
    odd_fill = PatternFill(fill_type="solid", fgColor="F7FBFF")
    even_fill = PatternFill(fill_type="solid", fgColor="EAF2F8")

    for col_index, header in enumerate(HEADERS, start=1):
        cell = ws.cell(row=1, column=col_index, value=header)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for row_index, row in enumerate(ROWS, start=2):
        fill = odd_fill if row_index % 2 == 0 else even_fill
        values = [
            row["Test Case ID"],
            VERSION,
            row["Feature"],
            row["Test Case"],
            row["Preconditions"],
            row["Steps to Execute"],
            row["Test Data"],
            row["Expected Output"],
            "TBD",
            "PENDING",
            "-",
            "-",
            "-",
            "-",
        ]
        for col_index, value in enumerate(values, start=1):
            cell = ws.cell(row=row_index, column=col_index, value=value)
            cell.fill = fill
            cell.alignment = Alignment(vertical="top", wrap_text=True)

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    autofit_columns(ws)
    wb.save(OUTPUT_PATH)


if __name__ == "__main__":
    main()
