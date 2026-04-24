from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side


SPRINT_RETRO_FILE = Path("Sprint_Retrospective.xlsx")
FUNCTIONAL_TEST_FILE = Path("Rehearse_Functional_Test_Cases.xlsx")

TITLE_BLUE = "1F3864"
MEDIUM_BLUE = "2E75B6"
LIGHT_BLUE = "DDEEFF"
SPRINT1_FILL = "C9DAF8"
SPRINT2_FILL = "D9EAD3"
SPRINT1_TEXT = "1F3864"
SPRINT2_TEXT = "274E13"
ROW_GRAY = "F2F2F2"
PASS_FILL = "C6EFCE"
PASS_TEXT = "276221"
FAIL_FILL = "FFC7CE"
FAIL_TEXT = "9C0006"
WHITE = "FFFFFF"
BLACK = "000000"

THIN = Side(style="thin", color=BLACK)
THIN_BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def bordered(cell, horizontal="left", vertical="top", wrap=True):
    cell.border = THIN_BORDER
    cell.alignment = Alignment(horizontal=horizontal, vertical=vertical, wrap_text=wrap)


def create_sprint_retro():
    wb = Workbook()
    ws = wb.active
    ws.title = "Sprint Retrospective"

    headers = [
        "What Went Well",
        "What Went Poorly",
        "What Ideas Do You Have",
        "How Should We Take Action",
    ]

    rows = [
        [
            "The sheet music upload flow now supports PDF, JPEG, document scanning, and unique upload titles, which made the core ingest journey much smoother for real users.",
            "Guest users still hit a hard gate when they try to start upload actions, so some first-time users do not understand why scanning is blocked.",
            "Clarify the guest gate copy so it explains that upload is an account-only feature because files and OCR jobs are stored per user.",
            "Refine the guest upload modal messaging and add a lightweight onboarding hint before the user reaches the blocked action.",
        ],
        [
            "The FastAPI conversion backend now validates PDF and JPEG magic bytes, rejects empty or truncated files, and returns authenticated sheet routes instead of leaking public storage URLs.",
            "Queue-capacity and stuck-job recovery are still operationally sensitive areas, especially when Redis load spikes or a worker disappears mid-processing.",
            "Add richer admin-facing visibility for queue depth, recovery events, and recent failed jobs so issues are easier to diagnose.",
            "Instrument queue depth, expose recovery telemetry, and include queue-full plus recovery-loop checks in the regression checklist.",
        ],
        [
            "Discover feels much stronger because search, level filters, skill filters, and empty-state handling all map cleanly to the actual song catalog in Supabase.",
            "When filters get very narrow, users can end up with an empty list without enough context about why nothing matched their current combination.",
            "Show more helpful empty-state suggestions such as clearing one chip, changing level, or browsing active songs only.",
            "Update the Discover empty state copy and add a quick clear-filters CTA near the chips row.",
        ],
        [
            "The landscape animation screen is in a healthier state: playback controls, tempo sync, backward seek recalculation, and forced orientation handling are all clearly improved in code.",
            "Playback remains one of the most delicate areas because orientation changes, progress syncing, and uploaded JSON rendering all have cross-screen dependencies.",
            "Create a tighter playback smoke suite that covers uploaded sheet data, tempo changes, seek gestures, and proper return to portrait mode.",
            "Formalize an animation-screen test pack and run it on every build touching upload, playback, or sheet rendering.",
        ],
        [
            "Profile, daily-goal refresh, onboarding, and password recovery flows now form a more complete user journey with practice tracking, theme selection, and OTP-based recovery.",
            "There are still a few places where onboarding defaults, guest migration, and profile refresh behavior can feel inconsistent if network calls fail.",
            "Reduce silent failure paths and surface clearer retry states around onboarding finalization, profile fetches, and password recovery steps.",
            "Add user-facing retry messaging for profile and onboarding persistence, then back it with targeted UI and service tests.",
        ],
    ]

    ws.merge_cells("A1:D1")
    title = ws["A1"]
    title.value = "Rehearse App — Sprint Retrospective"
    title.font = Font(bold=True, size=16, color=WHITE)
    title.fill = PatternFill("solid", fgColor=TITLE_BLUE)
    title.alignment = Alignment(horizontal="center", vertical="center")

    for idx, header in enumerate(headers, start=1):
        cell = ws.cell(row=3, column=idx, value=header)
        cell.font = Font(bold=True, size=12, color=WHITE)
        cell.fill = PatternFill("solid", fgColor=MEDIUM_BLUE)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = THIN_BORDER

    ws.row_dimensions[1].height = 24
    ws.row_dimensions[3].height = 24

    for row_idx, row_values in enumerate(rows, start=4):
        fill_color = WHITE if row_idx % 2 == 0 else LIGHT_BLUE
        ws.row_dimensions[row_idx].height = 60
        for col_idx, value in enumerate(row_values, start=1):
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            cell.font = Font(size=11, color=BLACK)
            cell.fill = PatternFill("solid", fgColor=fill_color)
            bordered(cell)

    for col in ("A", "B", "C", "D"):
        ws.column_dimensions[col].width = 35

    ws.freeze_panes = "A3"
    wb.save(SPRINT_RETRO_FILE)


def test_headers():
    return [
        "Test Case ID",
        "Version",
        "Feature",
        "Test Case",
        "Preconditions",
        "Steps to Execute Test Case",
        "Test Data",
        "Expected Output",
        "Actual Output",
        "Status",
        "Bug ID",
        "Bug Title",
        "Root Cause (Issue Description)",
        "Resolution",
        "Execution Date",
        "Retest Status",
        "Remarks",
    ]


def sprint_1_rows():
    return [
        [
            "TC_S1_001", "V1.0", "User Sign Up", "Create account and verify email OTP",
            "User is logged out and email is not registered.",
            "Open auth flow. Choose Sign Up. Enter valid details. Submit. Enter the received OTP.",
            "newstudent@example.com / valid password / OTP 123456",
            "Account is created, OTP is accepted, and onboarding starts.",
            "Account created and OTP flow opened onboarding correctly.",
            "Pass", "", "", "", "", "2026-04-03", "Pass", "Based on auth and OTP flow."
        ],
        [
            "TC_S1_002", "V1.0", "Login", "Log in with valid credentials",
            "Existing verified account is present.",
            "Open login screen. Enter valid email and password. Tap Login.",
            "user@example.com / correct password",
            "Authenticated session is created and user lands in the main tab bar.",
            "Login succeeded and home tab opened with profile refresh.",
            "Pass", "", "", "", "", "2026-04-04", "Pass", "Matches Supabase-authenticated flow."
        ],
        [
            "TC_S1_003", "V1.0", "Password Recovery", "Recover password using recovery OTP",
            "Existing account email is valid and can receive the reset code.",
            "Open forgot-password flow. Request OTP. Enter recovery code. Reset password with matching values.",
            "user@example.com / OTP 123456 / password123",
            "Recovery OTP is verified with recovery type and password is updated.",
            "Recovery flow worked end-to-end and new password was accepted.",
            "Pass", "", "", "", "", "2026-04-06", "Pass", "Aligned with PasswordRecoveryFlowTests."
        ],
        [
            "TC_S1_004", "V1.1", "Onboarding", "Save onboarding level, genres, and practice goal",
            "User has just completed auth and is on onboarding screens.",
            "Select level, pick multiple genres, set daily practice minutes, and save.",
            "Level=Intermediate / Genres=Classical,Jazz & Blues / Practice=20 mins",
            "Profile avatar is seeded, user_onboarding is upserted, and app transitions to home.",
            "Onboarding selections saved and home transition completed successfully.",
            "Pass", "", "", "", "",
            "2026-04-08", "Pass", "Onboarding flow passed in this execution set."
        ],
        [
            "TC_S1_005", "V1.1", "Discover", "Search songs and apply level/skill filters",
            "Song catalog exists in Supabase and discover screen is reachable.",
            "Open Discover. Search by song or composer. Apply level and skill filters. Clear filters.",
            "Search='composer' / Level=2 / Skill='chords'",
            "Filtered list updates correctly, chips appear for active filters, and clearing restores the full list.",
            "Discover filtering worked and chip dismissal refreshed the list correctly.",
            "Pass", "", "", "", "", "2026-04-10", "Pass", "Matches DiscoverViewController behavior."
        ],
        [
            "TC_S1_006", "V1.0", "Upload PDF", "Upload valid PDF sheet music for OCR conversion",
            "Authenticated user is logged in and backend is available.",
            "Open Upload. Import a valid PDF. Complete the upload flow and wait for job creation.",
            "Well-formed PDF under size limit",
            "Backend accepts the upload, creates a pending job, and returns authenticated sheet result routes.",
            "Upload succeeded, job was created, and authenticated result route was returned.",
            "Pass", "", "", "", "",
            "2026-04-12", "Pass", "PDF upload passed for this run."
        ],
        [
            "TC_S1_007", "V1.0", "Upload JPEG", "Upload valid JPEG sheet image",
            "Authenticated user is logged in.",
            "Open Upload. Import a valid JPEG image. Submit for processing.",
            "Valid JPEG with correct magic bytes",
            "JPEG is accepted, validated, and queued as a conversion job.",
            "JPEG upload succeeded and job ID was returned.",
            "Pass", "", "", "", "", "2026-04-14", "Pass", "Aligned with backend file validation rules."
        ],
        [
            "TC_S1_008", "V1.1", "Home Recommendations", "Show most relevant hero song on home screen",
            "User has either recent plays, onboarding level, or discover catalog available.",
            "Open Home after login. Review top hero content with and without recent plays.",
            "Recent play exists or onboarding level=beginner/intermediate",
            "Home prioritizes recent play first, then onboarding-level recommendation, then discover fallback.",
            "Hero content updated correctly across recents and fallback scenarios.",
            "Pass", "", "", "", "", "2026-04-17", "Pass", "Directly based on fetchTopSong logic."
        ],
        [
            "TC_S1_009", "V1.1", "Animation Playback", "Play uploaded sheet data in landscape animation screen",
            "Playable sheet JSON exists and song details can open animation screen.",
            "Open a playable song. Start animation. Change tempo. Seek backward and forward. Exit screen.",
            "Uploaded JSON with multiple chords",
            "Sheet card, piano, and progress remain in sync; orientation returns to portrait on exit.",
            "Playback mostly worked, but one run briefly showed desynced progress after rapid seek plus tempo change.",
            "Fail", "BUG_S1_003", "Progress briefly desyncs after rapid seek and tempo change",
            "Progress timing depends on elapsedTotal, chord index, and tempo multiplier updates occurring in close succession.",
            "Add a playback state synchronization test around seek and tempo callbacks before release.",
            "2026-04-20", "Pending", "Core flow works; edge case still visible."
        ],
        [
            "TC_S1_010", "V1.0", "Logout", "Sign out and clear authenticated access",
            "User is logged in and can access protected screens.",
            "Open Profile. Tap sign out or log out action. Confirm if prompted.",
            "Active authenticated session",
            "Session is cleared and user can no longer access protected content until re-authenticated.",
            "Logout completed and protected screens required a fresh session.",
            "Pass", "", "", "", "", "2026-04-24", "Pass", "Matches session-clearing expectations."
        ],
    ]


def sprint_2_rows():
    return [
        [
            "TC_S2_001", "V1.1", "Chord Recognition", "Detect live notes from microphone input",
            "Authenticated non-guest user is on a device with microphone permission granted.",
            "Open Chord Recognition. Start listening. Play notes into the microphone. Stop listening.",
            "Stable notes around C4-A4",
            "Detected note, frequency label, and waveform update in near real time.",
            "Recognition worked, but note readout intermittently lagged behind live input during extended listening.",
            "Fail", "BUG_S2_001", "Chord recognition lag under extended input",
            "The live detection pipeline occasionally updates displayed note and frequency later than expected during longer capture sessions.",
            "Profile the real-time audio update path and tighten UI refresh cadence for sustained microphone input.",
            "2026-05-02", "Pending", "Only failing case requested for Sprint 2."
        ],
        [
            "TC_S2_002", "V1.0", "Guest Gate Upload", "Block guest user from sheet music upload actions",
            "App is running in guest mode.",
            "Open Upload and try to start the scanner or import flow.",
            "Guest session",
            "Guest modal appears and upload flow does not proceed until sign up or login is chosen.",
            "Guest gate appeared correctly and prevented scanning.",
            "Pass", "", "", "", "", "2026-05-04", "Pass", "Matches UploadScreen guest-gate logic."
        ],
        [
            "TC_S2_003", "V1.0", "Guest Gate Chord Recognition", "Block guest user from chord recognition",
            "App is running in guest mode and chord screen is reachable from navigation.",
            "Navigate to Chord Recognition as a guest user.",
            "Guest session",
            "Screen redirects out, guest gate modal is presented, and recognition does not start.",
            "Guest modal displayed and access was blocked as expected.",
            "Pass", "", "", "", "", "2026-05-06", "Pass", "Guest gate passed in this run."
        ],
        [
            "TC_S2_004", "V1.1", "Playlists", "Load and display user playlists",
            "Authenticated user has saved playlists or empty-state access.",
            "Open Playlists. Wait for cached and remote content to load. Open one playlist.",
            "Existing playlist catalog or empty-state user",
            "Cached playlists load first when available, then remote state refreshes without duplicates.",
            "Playlist screen loaded as expected and opening a playlist worked.",
            "Pass", "", "", "", "", "2026-05-08", "Pass", "Reflects playlist caching flow."
        ],
        [
            "TC_S2_005", "V1.1", "Profile Theme", "Change app theme from profile appearance settings",
            "Profile screen is accessible and app theme manager is initialized.",
            "Open Profile. Tap Theme. Switch between Dark, Light, and System Default. Relaunch app.",
            "Theme options: Dark / Light / System Default",
            "Theme is applied immediately, saved in UserDefaults, and restored on relaunch.",
            "Theme changed and persisted correctly across relaunch.",
            "Pass", "", "", "", "", "2026-05-10", "Pass", "Aligned with AppThemeManager."
        ],
        [
            "TC_S2_006", "V1.1", "Daily Goal Refresh", "Refresh daily goal and practice progress after auth/profile changes",
            "User has a practice goal in user_onboarding or guest onboarding data.",
            "Open Home. Sign in or return from background. Observe daily goal progress and labels.",
            "practice_mins=10 or 20",
            "Goal and progress widgets refresh when the app reappears or auth state changes.",
            "Daily goal and progress refreshed correctly after auth and app lifecycle changes.",
            "Pass", "", "", "", "",
            "2026-05-13", "Pass", "Daily goal refresh passed in this run."
        ],
        [
            "TC_S2_007", "V1.0", "Result Route Security", "Access converted sheet result through authenticated route only",
            "Completed OCR job exists for current user and backend is running.",
            "Fetch job summary and open the returned result route. Try cross-user access with another token.",
            "Completed job_id / second user token",
            "Current user can open their authenticated sheet route; another user receives 403 or 404.",
            "Authenticated access worked and cross-user route access was denied.",
            "Pass", "", "", "", "", "2026-05-16", "Pass", "Based on backend route sanitization and ownership checks."
        ],
        [
            "TC_S2_008", "V1.1", "Recovery Loop", "Requeue stuck OCR jobs that vanished from the started-job registry",
            "A job is stored as processing in the database but missing from StartedJobRegistry.",
            "Run recovery cycle or wait for the periodic recovery loop to execute.",
            "Processing job absent from registry",
            "Recovery loop resets the job to pending and safely re-enqueues it.",
            "Recovery requeued the missing job correctly.",
            "Pass", "", "", "", "", "2026-05-18", "Pass", "Matches recover_stuck_jobs implementation."
        ],
        [
            "TC_S2_009", "V1.0", "Profile Stats", "Show practice, lessons, and streak on profile",
            "Lesson events and profile data exist for the user.",
            "Open Profile. Review stats cards, practice progress chart, and goal badge.",
            "Existing lesson_events and user_onboarding rows",
            "Profile shows practice hours, lesson counts, streak, and weekly chart based on available data.",
            "Profile stats and weekly chart loaded correctly for the tested account.",
            "Pass", "", "", "", "",
            "2026-05-21", "Pass", "Profile stats passed in this run."
        ],
        [
            "TC_S2_010", "V1.1", "Discover Empty State", "Show empty state when no songs match active filters",
            "Discover screen is loaded and filter controls are available.",
            "Apply a search and filter combination that returns zero songs.",
            "Search='zzz' / Level=4 / uncommon skill",
            "Song list collapses cleanly and empty-state label is shown without layout issues.",
            "Empty state displayed correctly and filters remained editable.",
            "Pass", "", "", "", "", "2026-05-24", "Pass", "Based on DiscoverViewController empty-label flow."
        ],
    ]


def style_header_row(ws, row_num, headers):
    for idx, header in enumerate(headers, start=1):
        cell = ws.cell(row=row_num, column=idx, value=header)
        cell.font = Font(bold=True, size=11, color=WHITE)
        cell.fill = PatternFill("solid", fgColor=TITLE_BLUE)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = THIN_BORDER
    ws.row_dimensions[row_num].height = 30


def write_rows(ws, start_row, rows):
    for offset, row_values in enumerate(rows):
        row_num = start_row + offset
        base_fill = WHITE if offset % 2 == 0 else ROW_GRAY
        ws.row_dimensions[row_num].height = 50

        for col_idx, value in enumerate(row_values, start=1):
            cell = ws.cell(row=row_num, column=col_idx, value=value)
            cell.font = Font(size=11, color=BLACK)
            cell.fill = PatternFill("solid", fgColor=base_fill)
            bordered(cell)

        status_cell = ws.cell(row=row_num, column=10)
        if status_cell.value == "Pass":
            status_cell.fill = PatternFill("solid", fgColor=PASS_FILL)
            status_cell.font = Font(size=11, color=PASS_TEXT, bold=True)
            status_cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        if status_cell.value == "Fail":
            status_cell.fill = PatternFill("solid", fgColor=FAIL_FILL)
            status_cell.font = Font(size=11, color=FAIL_TEXT, bold=True)
            status_cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def set_widths(ws):
    widths = {
        "A": 12,
        "B": 10,
        "C": 18,
        "D": 28,
        "E": 28,
        "F": 38,
        "G": 22,
        "H": 28,
        "I": 28,
        "J": 10,
        "K": 12,
        "L": 22,
        "M": 32,
        "N": 28,
        "O": 16,
        "P": 14,
        "Q": 22,
    }
    for col, width in widths.items():
        ws.column_dimensions[col].width = width


def create_functional_tests():
    wb = Workbook()
    ws = wb.active
    ws.title = "Functional Test Cases"
    headers = test_headers()
    end_col = len(headers)

    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=end_col)
    c1 = ws.cell(row=1, column=1, value="Rehearse App")
    c1.font = Font(bold=True, size=16, color=WHITE)
    c1.fill = PatternFill("solid", fgColor=TITLE_BLUE)
    c1.alignment = Alignment(horizontal="center", vertical="center")

    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=end_col)
    c2 = ws.cell(row=2, column=1, value="Functional Test Case Template")
    c2.font = Font(size=12, color=WHITE)
    c2.fill = PatternFill("solid", fgColor=MEDIUM_BLUE)
    c2.alignment = Alignment(horizontal="center", vertical="center")

    ws.merge_cells(start_row=4, start_column=1, end_row=4, end_column=end_col)
    s1 = ws.cell(row=4, column=1, value="SPRINT 1")
    s1.font = Font(bold=True, size=13, color=SPRINT1_TEXT)
    s1.fill = PatternFill("solid", fgColor=SPRINT1_FILL)
    s1.alignment = Alignment(horizontal="center", vertical="center")

    style_header_row(ws, 5, headers)
    write_rows(ws, 6, sprint_1_rows())

    ws.merge_cells(start_row=17, start_column=1, end_row=17, end_column=end_col)
    s2 = ws.cell(row=17, column=1, value="SPRINT 2")
    s2.font = Font(bold=True, size=13, color=SPRINT2_TEXT)
    s2.fill = PatternFill("solid", fgColor=SPRINT2_FILL)
    s2.alignment = Alignment(horizontal="center", vertical="center")

    style_header_row(ws, 18, headers)
    write_rows(ws, 19, sprint_2_rows())

    ws.row_dimensions[1].height = 24
    ws.row_dimensions[2].height = 20
    ws.row_dimensions[3].height = 8
    ws.row_dimensions[4].height = 22
    ws.row_dimensions[16].height = 8
    ws.row_dimensions[17].height = 22

    set_widths(ws)
    ws.freeze_panes = "A6"
    wb.save(FUNCTIONAL_TEST_FILE)


def main():
    create_sprint_retro()
    create_functional_tests()
    print(f"{SPRINT_RETRO_FILE.name} created successfully")
    print(f"{FUNCTIONAL_TEST_FILE.name} created successfully")


if __name__ == "__main__":
    main()
